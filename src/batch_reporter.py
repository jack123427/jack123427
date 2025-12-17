import os
import sys
import pandas as pd
from docx import Document
from docx.shared import Inches
import matplotlib.pyplot as plt
from datetime import datetime
import io

# This block allows the script to be run directly without import errors
# by adding the project's root directory to the Python path.
_root_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
if _root_dir not in sys.path:
    sys.path.insert(0, _root_dir)

from src.data_fetcher import get_top_100_stocks, get_historical_data
from src.model import train_and_predict

def create_prediction_chart(historical_df, forecast_df, stock_code, stock_name):
    """
    Generates a prediction chart and returns it as a bytes object.
    """
    plt.figure(figsize=(10, 6))

    history_to_plot = historical_df.tail(90)
    plt.plot(history_to_plot['Date'], history_to_plot['ClosingPrice'], label='歷史收盤價', color='blue')

    last_history_date = history_to_plot['Date'].iloc[-1]
    last_history_price = history_to_plot['ClosingPrice'].iloc[-1]

    forecast_plot_df = pd.concat([
        pd.DataFrame({'Date': [last_history_date], 'PredictedPrice': [last_history_price]}),
        forecast_df
    ])

    plt.plot(forecast_plot_df['Date'], forecast_plot_df['PredictedPrice'], label='預測收盤價', color='red', linestyle='--')

    plt.title(f'{stock_code} - {stock_name} 股價預測')
    plt.xlabel('日期')
    plt.ylabel('收盤價 (TWD)')
    plt.legend()
    plt.grid(True)
    plt.tight_layout()

    buf = io.BytesIO()
    plt.savefig(buf, format='png')
    plt.close()
    buf.seek(0)
    return buf

def main():
    """
    Main function to generate the batch report.
    """
    print("--- 開始產生百大個股預測報告 ---")

    print("步驟 1: 正在獲取市值百大公司列表...")
    top_100 = get_top_100_stocks()
    if top_100 is None or top_100.empty:
        print("錯誤：無法獲取百大公司列表，報告無法產生。")
        return

    document = Document()
    document.add_heading('台灣市值百大個股預測報告', 0)

    today_str = datetime.now().strftime('%Y-%m-%d')
    document.add_paragraph(f"報告產生日期：{today_str}")

    for index, row in top_100.iterrows():
        stock_code = row['公司代號']
        stock_name = row['公司簡稱']

        print(f"\n--- 正在處理 {index + 1}/{len(top_100)}: {stock_code} {stock_name} ---")

        try:
            print("  正在獲取歷史資料...")
            historical_df = get_historical_data([stock_code], period="1y")
            if historical_df is None or historical_df.empty:
                print(f"  警告：找不到 {stock_code} 的歷史資料，跳過。")
                continue

            print("  正在進行模型訓練與預測...")
            forecast_df, _ = train_and_predict(historical_df, forecast_days=30)
            if forecast_df is None or forecast_df.empty:
                print(f"  警告：為 {stock_code} 進行預測時發生錯誤，跳過。")
                continue

            print("  正在產生圖表...")
            chart_buffer = create_prediction_chart(historical_df, forecast_df, stock_code, stock_name)

            document.add_heading(f"{stock_code} {stock_name}", level=1)
            document.add_picture(chart_buffer, width=Inches(6.0))

            print(f"  成功處理 {stock_code} {stock_name}。")

        except Exception as e:
            print(f"  處理 {stock_code} 時發生未預期的錯誤: {e}")
            continue

    filename = f"百大個股預測_{today_str}.docx"
    document.save(filename)
    print(f"\n報告已成功儲存至：{os.path.abspath(filename)}")


if __name__ == '__main__':
    import matplotlib
    matplotlib.use('Agg')
    main()
